from __future__ import annotations

import base64
import json
import os
import stat
import zipfile
from collections.abc import Generator, Iterator
from contextlib import closing
from io import BytesIO
from pathlib import Path
from typing import Any

from deepagents.backends.protocol import FileData, ReadResult
from deepagents.backends.utils import slice_read_response

from skail.runtime.redaction import RedactionRegistry

DOCUMENT_SUFFIXES = frozenset({".pdf", ".docx", ".xlsx", ".pptx", ".ipynb"})
OFFICE_SUFFIXES = frozenset({".docx", ".xlsx", ".pptx"})
MAX_DOCUMENT_BYTES = 8 * 1024 * 1024
MAX_OFFICE_UNCOMPRESSED_BYTES = 32 * 1024 * 1024
MAX_OFFICE_MEMBERS = 1000
MAX_EXTRACTED_CHARS = 1_000_000
MAX_READ_LINES = 40
MAX_LINE_CHARS = 400
MAX_PDF_PAGES = 200
MAX_PDF_STREAM_BYTES = 8 * 1024 * 1024
MAX_IMAGE_BYTES = 8 * 1024 * 1024
MAX_XLSX_SHEETS = 50
MAX_XLSX_ROWS = 2000
MAX_XLSX_COLUMNS = 256


class DocumentReadError(ValueError):
    pass


def read_document(
    path: Path, *, offset: int, limit: int, redactor: RedactionRegistry
) -> ReadResult:
    """Extract a bounded text view while keeping document bytes out of model context."""
    if limit <= 0:
        return ReadResult(file_data=FileData(content="", encoding="utf-8"), no_lines_requested=True)
    try:
        raw = _read_bounded(path)
        if path.suffix.lower() in OFFICE_SUFFIXES:
            _check_office_archive(raw)
        with closing(_extract_lines(path.suffix.lower(), raw)) as source:
            parts: list[str] = []
            total = 0
            for part in source:
                total += len(part) + 1
                if total > MAX_EXTRACTED_CHARS:
                    raise DocumentReadError("document text exceeds the read_file extraction limit")
                parts.append(part)
        text = redactor.scrub_text("\n".join(parts) or "[Document has no extractable text]")
        wrapped = "\n".join(_wrapped_lines(text))
        return slice_read_response(
            FileData(content=wrapped, encoding="utf-8"),
            max(offset, 0),
            min(limit, MAX_READ_LINES),
        )
    except DocumentReadError as error:
        return ReadResult(error=str(error))
    except Exception:
        # Parser exceptions can contain document text or file paths; keep tool output bounded.
        return ReadResult(error="Could not extract text from this document")


def read_bom_text(
    path: Path, *, offset: int, limit: int, redactor: RedactionRegistry
) -> ReadResult:
    if limit <= 0:
        return ReadResult(file_data=FileData(content="", encoding="utf-8"), no_lines_requested=True)
    try:
        raw = _read_bounded(path)
        encoding = "utf-16" if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
        text = redactor.scrub_text(raw.decode(encoding))
        return slice_read_response(FileData(content=text, encoding="utf-8"), offset, limit)
    except DocumentReadError as error:
        return ReadResult(error=str(error))
    except (UnicodeError, OSError):
        return ReadResult(error="Could not decode this text file")


def read_image(path: Path) -> ReadResult:
    try:
        raw = _read_bounded(path, max_bytes=MAX_IMAGE_BYTES, kind="Image")
    except DocumentReadError as error:
        return ReadResult(error=str(error))
    except OSError:
        return ReadResult(error="Could not read this image safely")
    return ReadResult(
        file_data=FileData(content=base64.b64encode(raw).decode("ascii"), encoding="base64")
    )


def _read_bounded(
    path: Path, *, max_bytes: int | None = None, kind: str = "Document"
) -> bytes:
    if max_bytes is None:
        max_bytes = MAX_DOCUMENT_BYTES
    before = path.stat()
    if not stat.S_ISREG(before.st_mode):
        raise DocumentReadError("read_file requires a regular file")
    if before.st_size > max_bytes:
        raise DocumentReadError(f"{kind} exceeds the read_file size limit")
    descriptor = os.open(
        path, os.O_RDONLY | getattr(os, "O_BINARY", 0) | getattr(os, "O_NOFOLLOW", 0)
    )
    with os.fdopen(descriptor, "rb") as source:
        opened = os.fstat(source.fileno())
        if not stat.S_ISREG(opened.st_mode) or opened.st_size > max_bytes:
            raise DocumentReadError(f"{kind} exceeds the read_file size limit")
        raw = source.read(max_bytes + 1)
    after = path.stat()
    identity = (before.st_dev, before.st_ino, before.st_size, before.st_mtime_ns)
    if (
        len(raw) > max_bytes
        or identity != (opened.st_dev, opened.st_ino, opened.st_size, opened.st_mtime_ns)
        or identity != (after.st_dev, after.st_ino, after.st_size, after.st_mtime_ns)
    ):
        raise DocumentReadError(f"{kind} changed during read or exceeds the size limit")
    return raw


def _check_office_archive(raw: bytes) -> None:
    try:
        with zipfile.ZipFile(BytesIO(raw)) as archive:
            members = archive.infolist()
            if len(members) > MAX_OFFICE_MEMBERS or sum(
                member.file_size for member in members
            ) > MAX_OFFICE_UNCOMPRESSED_BYTES:
                raise DocumentReadError("Office document exceeds the extraction size limit")
            if any(member.flag_bits & 1 for member in members):
                raise DocumentReadError("encrypted Office documents are not supported")
    except zipfile.BadZipFile as error:
        raise DocumentReadError("Office document is not a valid archive") from error


def _extract_lines(suffix: str, raw: bytes) -> Generator[str, None, None]:
    if suffix == ".pdf":
        yield from _pdf_lines(raw)
    elif suffix == ".docx":
        yield from _docx_lines(raw)
    elif suffix == ".xlsx":
        yield from _xlsx_lines(raw)
    elif suffix == ".pptx":
        yield from _pptx_lines(raw)
    elif suffix == ".ipynb":
        yield from _notebook_lines(raw)
    else:
        raise DocumentReadError("document format is not supported")


def _pdf_lines(raw: bytes) -> Iterator[str]:
    from pypdf import PdfReader, apply_configuration
    from pypdf.errors import LimitReachedError

    with apply_configuration(
        maximum_declared_stream_length=MAX_PDF_STREAM_BYTES,
        array_based_stream_maximum_output_length=MAX_PDF_STREAM_BYTES,
        lzw_maximum_output_length=MAX_PDF_STREAM_BYTES,
        run_length_maximum_output_length=MAX_PDF_STREAM_BYTES,
        zlib_maximum_output_length=MAX_PDF_STREAM_BYTES,
    ):
        try:
            reader = PdfReader(BytesIO(raw), strict=False)
            if reader.is_encrypted:
                raise DocumentReadError("PDF is encrypted; provide an unlocked copy")
            page_count = len(reader.pages)
            for index in range(min(page_count, MAX_PDF_PAGES)):
                page = reader.pages[index]
                contents = page.get_contents()
                if contents is not None and len(contents.get_data()) > MAX_PDF_STREAM_BYTES:
                    raise DocumentReadError("PDF decoded content exceeds the read_file limit")
                text = page.extract_text() or ""
                if text.strip():
                    yield f"[Page {index + 1}]"
                    yield text
                else:
                    yield f"[Page {index + 1}: no extractable text; OCR required]"
            if page_count > MAX_PDF_PAGES:
                yield f"[Remaining {page_count - MAX_PDF_PAGES} pages omitted: page limit reached]"
        except LimitReachedError as error:
            raise DocumentReadError("PDF decoded content exceeds the read_file limit") from error


def _docx_lines(raw: bytes) -> Iterator[str]:
    from docx import Document
    from docx.table import Table

    document = Document(BytesIO(raw))
    table_number = 0
    for block in document.iter_inner_content():
        if isinstance(block, Table):
            table_number += 1
            yield f"[Table {table_number}]"
            for row in block.rows:
                yield " | ".join(cell.text.replace("\n", " / ") for cell in row.cells)
        elif block.text.strip():
            yield block.text


def _xlsx_lines(raw: bytes) -> Iterator[str]:
    from openpyxl import load_workbook  # type: ignore[import-untyped]

    workbook = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    try:
        for sheet in workbook.worksheets[:MAX_XLSX_SHEETS]:
            yield f"[Sheet {sheet.title}]"
            row_limit = min(sheet.max_row or MAX_XLSX_ROWS, MAX_XLSX_ROWS)
            column_limit = min(sheet.max_column or MAX_XLSX_COLUMNS, MAX_XLSX_COLUMNS)
            for row in sheet.iter_rows(max_row=row_limit, max_col=column_limit):
                for cell in row:
                    if cell.value is not None:
                        yield f"{cell.coordinate}: {cell.value}"
            if sheet.max_row and sheet.max_row > MAX_XLSX_ROWS:
                yield f"[Rows after {MAX_XLSX_ROWS} omitted; use a spreadsheet tool]"
            if sheet.max_column and sheet.max_column > MAX_XLSX_COLUMNS:
                yield f"[Columns after {MAX_XLSX_COLUMNS} omitted; use a spreadsheet tool]"
        if len(workbook.worksheets) > MAX_XLSX_SHEETS:
            yield f"[Sheets after {MAX_XLSX_SHEETS} omitted; use a spreadsheet tool]"
    finally:
        workbook.close()


def _pptx_lines(raw: bytes) -> Iterator[str]:
    from pptx import Presentation

    presentation = Presentation(BytesIO(raw))
    for slide_number, slide in enumerate(presentation.slides, start=1):
        yield f"[Slide {slide_number}]"
        table_number = 0
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text.strip():
                yield shape.text
            if shape.has_table:
                table_number += 1
                yield f"[Table {table_number}]"
                for row in shape.table.rows:
                    yield " | ".join(cell.text.replace("\n", " / ") for cell in row.cells)


def _notebook_lines(raw: bytes) -> Iterator[str]:
    notebook = json.loads(raw.decode("utf-8-sig"))
    if not isinstance(notebook, dict) or not isinstance(notebook.get("cells"), list):
        raise DocumentReadError("notebook has no cell list")
    for number, cell in enumerate(notebook["cells"], start=1):
        if not isinstance(cell, dict):
            continue
        kind = str(cell.get("cell_type", "unknown")).capitalize()
        yield f"[{kind} cell {number}]"
        source: Any = cell.get("source", "")
        if isinstance(source, list):
            source = "".join(part for part in source if isinstance(part, str))
        if isinstance(source, str) and source:
            yield source
        if cell.get("outputs"):
            yield "[Cell outputs omitted]"


def _wrapped_lines(text: str) -> Iterator[str]:
    for line in text.splitlines():
        if not line:
            yield ""
            continue
        for start in range(0, len(line), MAX_LINE_CHARS):
            yield line[start : start + MAX_LINE_CHARS]
