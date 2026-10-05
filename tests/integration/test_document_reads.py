from __future__ import annotations

import json
import zipfile
from pathlib import Path

import pytest
from deepagents.backends import FilesystemBackend
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, NameObject

from skail.runtime.redaction import RedactionRegistry
from skail.tools.backend import PolicyFilesystemBackend


def _backend(tmp_path: Path, redactor: RedactionRegistry | None = None) -> PolicyFilesystemBackend:
    return PolicyFilesystemBackend(
        tmp_path, redactor=redactor or RedactionRegistry(), task_id="document-test"
    )


def _text(backend: PolicyFilesystemBackend, name: str) -> str:
    result = backend.read(name)
    assert result.error is None
    assert result.file_data is not None
    assert result.file_data["encoding"] == "utf-8"
    return result.file_data["content"]


def test_pdf_read_is_paginated_and_labels_pages(tmp_path: Path) -> None:
    fixture = Path(__file__).parents[1] / "fixtures" / "documents" / "sample.pdf"
    (tmp_path / "sample.pdf").write_bytes(fixture.read_bytes())
    backend = _backend(tmp_path)

    first = backend.read("sample.pdf", offset=0, limit=2)
    second = backend.read("sample.pdf", offset=2, limit=2)

    assert first.error is None and first.file_data is not None
    assert second.error is None and second.file_data is not None
    assert first.file_data["encoding"] == second.file_data["encoding"] == "utf-8"
    assert "[Page 1]" in first.file_data["content"]
    assert "[Page 2]" in second.file_data["content"]
    assert first.next_offset == 2
    assert first.total_lines == second.total_lines


def test_image_only_pdf_reports_ocr_need(tmp_path: Path) -> None:
    fixture = Path(__file__).parents[1] / "fixtures" / "documents" / "blank.pdf"
    (tmp_path / "blank.pdf").write_bytes(fixture.read_bytes())

    content = _text(_backend(tmp_path), "blank.pdf")

    assert "OCR" in content
    assert "[Page 1" in content


def test_pdf_rejects_oversized_decoded_stream(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from skail.tools import document_reader

    writer = PdfWriter()
    page = writer.add_blank_page(width=200, height=200)
    stream = DecodedStreamObject()
    stream.set_data(b" " * (2 * 1024 * 1024))
    page[NameObject("/Contents")] = writer._add_object(stream.flate_encode())
    with (tmp_path / "compressed.pdf").open("wb") as output:
        writer.write(output)
    monkeypatch.setattr(document_reader, "MAX_PDF_STREAM_BYTES", 1024 * 1024, raising=False)

    result = _backend(tmp_path).read("compressed.pdf")

    assert result.error is not None
    assert "limit" in result.error.lower()


def test_docx_read_keeps_paragraphs_and_tables_and_redacts(tmp_path: Path) -> None:
    secret = "DOCUMENT-SECRET-CANARY"
    document = Document()
    document.add_paragraph(f"Project overview {secret}")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    document.save(tmp_path / "notes.docx")
    redactor = RedactionRegistry()
    redactor.register(secret)

    content = _text(_backend(tmp_path, redactor), "notes.docx")

    assert "Project overview" in content
    assert "[Table 1]" in content
    assert "Metric" in content and "Value" in content
    assert secret not in content


def test_xlsx_read_labels_sheets_cells_and_preserves_formulas(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Budget"
    sheet["A1"] = "Item"
    sheet["B1"] = 42
    sheet["A2"] = "Formula"
    sheet["B2"] = "=1+1"
    workbook.save(tmp_path / "budget.xlsx")

    content = _text(_backend(tmp_path), "budget.xlsx")

    assert "[Sheet Budget]" in content
    assert "A1: Item" in content
    assert "B1: 42" in content
    assert "B2: =1+1" in content


def test_pptx_read_labels_slides_and_tables(tmp_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    textbox.text = "Quarterly results"
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(5), Inches(1)).table
    table.cell(0, 0).text = "Revenue"
    table.cell(0, 1).text = "100"
    presentation.save(tmp_path / "results.pptx")

    content = _text(_backend(tmp_path), "results.pptx")

    assert "[Slide 1]" in content
    assert "Quarterly results" in content
    assert "Revenue" in content and "100" in content


def test_notebook_read_includes_sources_but_not_saved_outputs(tmp_path: Path) -> None:
    notebook = {
        "cells": [
            {"cell_type": "markdown", "source": ["# Analysis\n"]},
            {
                "cell_type": "code",
                "source": ["print('hello')\n"],
                "outputs": [{"output_type": "stream", "text": ["PRIVATE-OUTPUT\n"]}],
            },
        ],
        "nbformat": 4,
        "nbformat_minor": 5,
        "metadata": {},
    }
    (tmp_path / "analysis.ipynb").write_text(json.dumps(notebook), encoding="utf-8")

    content = _text(_backend(tmp_path), "analysis.ipynb")

    assert "[Markdown cell 1]" in content
    assert "[Code cell 2]" in content
    assert "print('hello')" in content
    assert "PRIVATE-OUTPUT" not in content
    assert "outputs omitted" in content.lower()


def test_utf16_text_read_is_decoded_without_base64(tmp_path: Path) -> None:
    (tmp_path / "notes.txt").write_text("Café status\nSecond line\n", encoding="utf-16")

    content = _text(_backend(tmp_path), "notes.txt")

    assert "Café status" in content
    assert "Second line" in content


def test_svg_is_read_as_text_instead_of_binary_image(tmp_path: Path) -> None:
    (tmp_path / "diagram.svg").write_text(
        '<svg xmlns="http://www.w3.org/2000/svg"><text>Label</text></svg>',
        encoding="utf-8",
    )

    content = _text(_backend(tmp_path), "diagram.svg")

    assert "<svg" in content
    assert "Label" in content


def test_image_read_does_not_use_unbounded_base_reader(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    (tmp_path / "pixel.png").write_bytes(b"\x89PNG\r\n\x1a\n")

    def reject_unbounded_read(*_args: object, **_kwargs: object) -> None:
        raise AssertionError("unbounded base reader used for image")

    monkeypatch.setattr(FilesystemBackend, "read", reject_unbounded_read)
    result = _backend(tmp_path).read("pixel.png")

    assert result.error is None
    assert result.file_data is not None
    assert result.file_data["encoding"] == "base64"


def test_unsupported_image_format_returns_actionable_error(tmp_path: Path) -> None:
    (tmp_path / "drawing.bmp").write_bytes(b"BM")

    result = _backend(tmp_path).read("drawing.bmp")

    assert result.file_data is None
    assert result.error is not None
    assert "PNG" in result.error


def test_archive_read_has_explicit_separate_tool_guidance(tmp_path: Path) -> None:
    with zipfile.ZipFile(tmp_path / "bundle.zip", "w") as archive:
        archive.writestr("inside.txt", "content")

    result = _backend(tmp_path).read("bundle.zip")

    assert result.file_data is None
    assert result.error is not None
    assert "archive" in result.error.lower()


def test_document_read_rejects_oversized_input_before_parsing(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from skail.tools import document_reader

    fixture = Path(__file__).parents[1] / "fixtures" / "documents" / "sample.pdf"
    (tmp_path / "sample.pdf").write_bytes(fixture.read_bytes())
    monkeypatch.setattr(document_reader, "MAX_DOCUMENT_BYTES", 100)

    result = _backend(tmp_path).read("sample.pdf")

    assert result.file_data is None
    assert result.error is not None
    assert "size limit" in result.error.lower()


def test_office_read_rejects_large_decompressed_archive(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from skail.tools import document_reader

    with zipfile.ZipFile(tmp_path / "large.docx", "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "x" * 1000)
    monkeypatch.setattr(document_reader, "MAX_OFFICE_UNCOMPRESSED_BYTES", 100)

    result = _backend(tmp_path).read("large.docx")

    assert result.file_data is None
    assert result.error is not None
    assert "size limit" in result.error.lower()
